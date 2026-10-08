import io
from datetime import datetime, date, timedelta
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse, HTMLResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from database import get_db
from database_models import Admin, Staff, Business
from utils.dependencies import get_current_admin
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/admin", tags=["Admin - Staff Reports"])

class ReportBusiness(BaseModel):
    business_id: int
    business_name: str
    category: str
    owner_name: str
    phone: str
    district: str
    created_at: datetime

class ReportStaffInfo(BaseModel):
    staff_id: str
    user_id: int
    name: str
    email: str
    phone: str
    state: Optional[str] = None
    district: Optional[str] = None

class ReportMeta(BaseModel):
    from_date: date
    to_date: date
    business_count: int

class StaffActivityReportResponse(BaseModel):
    staff: ReportStaffInfo
    report: ReportMeta
    businesses: List[ReportBusiness]

def get_staff_activity_data(staff_id: str, from_date: date, to_date: date, db: Session):
    if from_date > to_date:
        raise HTTPException(status_code=422, detail="from_date must not be greater than to_date")
        
    staff = db.query(Staff).filter(Staff.staff_id == staff_id).first()
    if not staff:
        raise HTTPException(status_code=404, detail=f"Field Staff with staff_id {staff_id} not found")
        
    if staff.role != "staff":
        raise HTTPException(status_code=403, detail="User is not a Field Staff")

    # The filter must be inclusive for from_date 00:00:00 to to_date 23:59:59
    next_day = to_date + timedelta(days=1)
    
    businesses = (
        db.query(Business)
        .filter(Business.staff_id == staff.staff_id)
        .filter(Business.created_at >= from_date)
        .filter(Business.created_at < next_day)
        .order_by(Business.created_at.desc())
        .all()
    )
    
    return staff, businesses

@router.get("/staff-report", response_model=StaffActivityReportResponse)
def get_staff_report(
    staff_id: str = Query(..., description="Staff ID to generate report for"),
    from_date: date = Query(..., description="Start date (YYYY-MM-DD)"),
    to_date: date = Query(..., description="End date (YYYY-MM-DD)"),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    staff, businesses = get_staff_activity_data(staff_id, from_date, to_date, db)
    
    report_businesses = [
        ReportBusiness(
            business_id=b.id,
            business_name=b.business_name,
            category=b.business_category,
            owner_name=b.owner_name,
            phone=b.owner_phone,
            district=b.district,
            created_at=b.created_at
        ) for b in businesses
    ]
    
    return StaffActivityReportResponse(
        staff=ReportStaffInfo(
            staff_id=staff.staff_id,
            user_id=staff.id,
            name=staff.name,
            email=staff.email,
            phone=staff.phone,
            state=staff.state,
            district=staff.district
        ),
        report=ReportMeta(
            from_date=from_date,
            to_date=to_date,
            business_count=len(businesses)
        ),
        businesses=report_businesses
    )

@router.get("/staff-report/export")
def export_staff_report(
    staff_id: str = Query(..., description="Staff ID to generate report for"),
    from_date: date = Query(..., description="Start date (YYYY-MM-DD)"),
    to_date: date = Query(..., description="End date (YYYY-MM-DD)"),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    staff, businesses = get_staff_activity_data(staff_id, from_date, to_date, db)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Activity Report"
    
    # Fonts and alignments
    title_font = Font(size=14, bold=True)
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="EEEEEE", end_color="EEEEEE", fill_type="solid")
    center_aligned_text = Alignment(horizontal="center")
    
    # Title
    ws.merge_cells('A1:C1')
    ws['A1'] = "FIELD STAFF ACTIVITY REPORT"
    ws['A1'].font = title_font
    ws['A1'].alignment = center_aligned_text
    
    ws.append([])
    
    # Staff Details
    ws.append(["Staff ID:", staff.staff_id])
    ws.append(["Staff Name:", staff.name])
    ws.append(["From Date:", from_date.strftime("%d-%m-%Y")])
    ws.append(["To Date:", to_date.strftime("%d-%m-%Y")])
    ws.append([])
    ws.append(["Total Businesses Added:", len(businesses)])
    ws.append([])
    
    # Table Header
    headers = ["Business Name", "Category", "Added Date", "Owner Name"]
    ws.append(headers)
    header_row = ws.max_row
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
    
    # Data Rows
    for b in businesses:
        date_str = b.created_at.strftime("%d-%m-%Y") if b.created_at else ""
        ws.append([b.business_name, b.business_category, date_str, b.owner_name])
        
    # Adjust column widths
    for col_idx in range(1, len(headers) + 1):
        column_letter = get_column_letter(col_idx)
        ws.column_dimensions[column_letter].width = 25
        
    # Freeze the table header
    ws.freeze_panes = f"A{header_row + 1}"
    
    # Export
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    filename = f"staff_report_{staff_id}_{from_date}_to_{to_date}.xlsx"
    headers = {
        'Content-Disposition': f'attachment; filename="{filename}"'
    }
    
    return StreamingResponse(stream, headers=headers, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# =============================================================================
# NEW NO-DATE APIS
# =============================================================================

no_date_router = APIRouter(prefix="/admin", tags=["Admin - Staff Reports (No Date)"])

class ReportMetaNoDate(BaseModel):
    business_count: int

class StaffActivityReportNoDateResponse(BaseModel):
    staff: ReportStaffInfo
    report: ReportMetaNoDate
    businesses: List[ReportBusiness]

def get_staff_all_activity_data(staff_id: str, db: Session):
    staff = db.query(Staff).filter(Staff.staff_id == staff_id).first()
    if not staff:
        raise HTTPException(status_code=404, detail=f"Field Staff with staff_id {staff_id} not found")
        
    if staff.role != "staff":
        raise HTTPException(status_code=403, detail="User is not a Field Staff")

    businesses = (
        db.query(Business)
        .filter(Business.staff_id == staff.staff_id)
        .order_by(Business.created_at.desc())
        .all()
    )
    
    return staff, businesses

@no_date_router.get("/staff-report/staff", response_model=StaffActivityReportNoDateResponse)
def get_staff_report_no_date(
    staff_id: str = Query(..., description="Staff ID to generate report for"),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    staff, businesses = get_staff_all_activity_data(staff_id, db)
    
    report_businesses = [
        ReportBusiness(
            business_id=b.id,
            business_name=b.business_name,
            category=b.business_category,
            owner_name=b.owner_name,
            phone=b.owner_phone,
            district=b.district,
            created_at=b.created_at
        ) for b in businesses
    ]
    
    return StaffActivityReportNoDateResponse(
        staff=ReportStaffInfo(
            staff_id=staff.staff_id,
            user_id=staff.id,
            name=staff.name,
            email=staff.email,
            phone=staff.phone,
            state=staff.state,
            district=staff.district
        ),
        report=ReportMetaNoDate(
            business_count=len(businesses)
        ),
        businesses=report_businesses
    )

@no_date_router.get("/staff-report/staff/export")
def export_staff_report_no_date(
    staff_id: str = Query(..., description="Staff ID to generate report for"),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    staff, businesses = get_staff_all_activity_data(staff_id, db)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Activity Report"
    
    # Fonts and alignments
    title_font = Font(size=14, bold=True)
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="EEEEEE", end_color="EEEEEE", fill_type="solid")
    center_aligned_text = Alignment(horizontal="center")
    
    # Title
    ws.merge_cells('A1:C1')
    ws['A1'] = "FIELD STAFF ACTIVITY REPORT"
    ws['A1'].font = title_font
    ws['A1'].alignment = center_aligned_text
    
    ws.append([])
    
    # Staff Details
    ws.append(["Staff ID:", staff.staff_id])
    ws.append(["Staff Name:", staff.name])
    ws.append([])
    ws.append(["Total Businesses Added:", len(businesses)])
    ws.append([])
    
    # Table Header
    headers = ["Business Name", "Category", "Added Date", "Owner Name"]
    ws.append(headers)
    header_row = ws.max_row
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
    
    # Data Rows
    for b in businesses:
        date_str = b.created_at.strftime("%d-%m-%Y") if b.created_at else ""
        ws.append([b.business_name, b.business_category, date_str, b.owner_name])
        
    # Adjust column widths
    for col_idx in range(1, len(headers) + 1):
        column_letter = get_column_letter(col_idx)
        ws.column_dimensions[column_letter].width = 25
        
    # Freeze the table header
    ws.freeze_panes = f"A{header_row + 1}"
    
    # Export
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    filename = f"staff_report_{staff_id}.xlsx"
    headers = {
        'Content-Disposition': f'attachment; filename="{filename}"'
    }
    
    return StreamingResponse(stream, headers=headers, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def generate_html_report(staff, businesses, from_date=None, to_date=None):
    date_html = ""
    if from_date and to_date:
        date_html = f"""
        <tr><td class="label">From Date:</td><td class="value">{from_date.strftime("%d-%m-%Y")}</td></tr>
        <tr><td class="label">To Date:</td><td class="value">{to_date.strftime("%d-%m-%Y")}</td></tr>
        """

    rows_html = ""
    for b in businesses:
        added_date = b.created_at.strftime("%d-%m-%Y") if b.created_at else ""
        rows_html += f"""
        <tr>
            <td>{b.business_name}</td>
            <td>{b.business_category}</td>
            <td>{added_date}</td>
            <td>{b.owner_name}</td>
        </tr>
        """

    html_content = f"""
    <!DOCTYPE html>
    <html lang="en">
    <head>
        <meta charset="UTF-8">
        <title>Field Staff Activity Report</title>
        <style>
            @page {{
                size: A4 portrait;
                margin: 15mm;
            }}
            body {{
                font-family: Arial, sans-serif;
                margin: 0;
                padding: 0;
                color: #000;
            }}
            .header {{
                text-align: center;
                margin-bottom: 20px;
            }}
            .header h1 {{
                margin: 0;
                font-size: 18px;
                text-transform: uppercase;
            }}
            .info-table {{
                margin-bottom: 20px;
                border: none;
                width: auto;
            }}
            .info-table th, .info-table td {{
                border: none;
                padding: 3px 15px 3px 0;
                text-align: left;
                font-size: 12px;
            }}
            .info-table .label {{
                font-weight: normal;
                white-space: nowrap;
            }}
            .info-table .value {{
                font-weight: normal;
            }}
            .data-table {{
                width: 100%;
                border-collapse: collapse;
            }}
            .data-table th, .data-table td {{
                border: 1px solid #000;
                padding: 6px;
                text-align: left;
                font-size: 11px;
            }}
            .data-table th {{
                background-color: #f2f2f2;
                font-weight: bold;
            }}
            thead {{
                display: table-header-group;
            }}
            tr {{
                page-break-inside: avoid;
            }}
            @media print {{
                body {{
                    -webkit-print-color-adjust: exact;
                    print-color-adjust: exact;
                }}
            }}
        </style>
    </head>
    <body>
        <div class="header">
            <h1>FIELD STAFF ACTIVITY REPORT</h1>
        </div>
        
        <table class="info-table">
            <tr><td class="label">Staff ID:</td><td class="value">{staff.staff_id}</td></tr>
            <tr><td class="label">Staff Name:</td><td class="value">{staff.name}</td></tr>
            {date_html}
            <tr><td class="label">Total Businesses Added:</td><td class="value">{len(businesses)}</td></tr>
        </table>

        <table class="data-table">
            <thead>
                <tr>
                    <th>Business Name</th>
                    <th>Category</th>
                    <th>Added Date</th>
                    <th>Owner Name</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>
    </body>
    </html>
    """
    return html_content

@router.get("/staff-report/print")
def print_staff_report(
    staff_id: str = Query(..., description="Staff ID to generate report for"),
    from_date: date = Query(..., description="Start date (YYYY-MM-DD)"),
    to_date: date = Query(..., description="End date (YYYY-MM-DD)"),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    staff, businesses = get_staff_activity_data(staff_id, from_date, to_date, db)
    html_content = generate_html_report(staff, businesses, from_date, to_date)
    return HTMLResponse(content=html_content)

@no_date_router.get("/staff-report/staff/print")
def print_staff_report_no_date(
    staff_id: str = Query(..., description="Staff ID to generate report for"),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    staff, businesses = get_staff_all_activity_data(staff_id, db)
    html_content = generate_html_report(staff, businesses)
    return HTMLResponse(content=html_content)
